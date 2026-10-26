import pytest
from docx import Document

from app.knowledge import (
    ConflictError,
    DocumentError,
    TextDocument,
    create_note,
    docx_html,
    list_documents,
    validate_title,
)


@pytest.mark.parametrize(
    "title",
    ["", " ", "../escape", "A/B", "A\\B", "CON", "con.txt", "LPT1", "trailing.", " leading", "x\x00y"],
)
def test_invalid_filenames(title):
    with pytest.raises(DocumentError):
        validate_title(title)


def test_note_is_real_utf8_and_never_overwrites(tmp_path):
    doc = create_note(tmp_path, "Gespräch ä 日本", "First\nSecond")
    assert doc.path.read_bytes() == "First\nSecond".encode()
    doc.save("Updated\nText")
    assert TextDocument.open(doc.path).text == "Updated\nText"
    with pytest.raises(DocumentError):
        create_note(tmp_path, "GESPRÄCH Ä 日本", "Overwrite")
    assert doc.path.read_text() == "Updated\nText"
    assert not list(tmp_path.glob(".getitdone-*"))


def test_external_changes_and_deleted_file_keep_editor_data(tmp_path):
    doc = create_note(tmp_path, "Note", "original")
    doc.path.write_text("external")
    with pytest.raises(ConflictError):
        doc.save("my edits")
    assert doc.path.read_text() == "external"
    assert doc.text == "original"
    doc.path.unlink()
    with pytest.raises(DocumentError):
        doc.save("my edits")
    assert not doc.path.exists()


def test_symlink_preserved(tmp_path):
    target = tmp_path / "target.txt"
    target.write_text("original")
    link = tmp_path / "link.txt"
    try:
        link.symlink_to(target)
    except OSError:
        pytest.skip("Symlink permission unavailable")
    doc = TextDocument.open(link)
    doc.save("changed")
    assert link.is_symlink() and target.read_text() == "changed"


def test_listing_formats_sort_refresh_and_missing_folder(tmp_path):
    for name in ("z.md", "A.txt", "b.docx", "C.pdf", "A.doc", "skip.exe"):
        (tmp_path / name).write_text("")
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub/hidden.txt").write_text("")
    assert [p.name for p in list_documents(tmp_path)] == ["A.doc", "A.txt", "b.docx", "C.pdf", "z.md"]
    (tmp_path / "A.txt").unlink()
    assert len(list_documents(tmp_path)) == 4
    with pytest.raises(DocumentError):
        list_documents(tmp_path / "missing")


def test_docx_text_tables_and_escaped_content(tmp_path):
    path = tmp_path / "meeting.docx"
    doc = Document()
    doc.add_paragraph("<script> & Erkenntnis")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Column A"
    table.cell(0, 1).text = "Column B"
    doc.save(path)
    preview = docx_html(path)
    assert "&lt;script&gt;" in preview and "Column A" in preview and "<table" in preview
    path.write_bytes(b"not a zip")
    with pytest.raises(DocumentError):
        docx_html(path)


def test_large_and_invalid_text_are_explicit_errors(tmp_path):
    path = tmp_path / "bad.txt"
    path.write_bytes(b"\x80\x81")
    with pytest.raises(DocumentError):
        TextDocument.open(path)
    path.write_bytes(b"x" * (8 * 1024 * 1024 + 1))
    with pytest.raises(DocumentError):
        TextDocument.open(path)
